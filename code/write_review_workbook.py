"""write_review_workbook.py: build the Excel review workbook (called by 02_candidates.py).

Plain-language note: makes data/curation/irAE-Map_review.xlsx, the file
Onyekachi works in. Yellow columns are for your decisions; dropdowns limit
them to valid values. You never need to open this script.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

YELLOW = PatternFill("solid", fgColor="FFF4C2")
GREY = PatternFill("solid", fgColor="E8E8E8")
BOLD = Font(bold=True)

ROLES = "accept,primary,supporting,lookalike,companion,exclude"
CONF = "High,Moderate,Low"

HOW_TO = [
    ("irAE-Map review workbook", True),
    ("", False),
    ("What this is: every candidate link between an irAE category and a CTCAE term, ICD-10-CM code, or FDA reaction term that the search rules found. Claude filled in a suggested role, confidence, and reason (grey columns). You make the decision (yellow columns). Nothing is final until you fill the yellow columns.", False),
    ("", False),
    ("How to decide a row (about 5 seconds for easy rows):", True),
    ("1. Read the term and the suggestion.", False),
    ("2. In author_role, pick one: accept (you agree with the suggested role, confidence AND reason exactly), or primary / supporting / lookalike / companion / exclude to set your own.", False),
    ("3. If you did not pick accept, also pick author_confidence and write author_rationale in your own words (one short sentence).", False),
    ("4. Optional: author_lookalike_note for diagnoses that are easily confused with this row.", False),
    ("5. Put your initials in decided_by and the date in decided_on (YYYY-MM-DD).", False),
    ("", False),
    ("Roles", True),
    ("primary: names the condition itself; enough on its own to identify the irAE category.", False),
    ("supporting: a lab, symptom, or grading term that fits the category but is not specific to it.", False),
    ("lookalike: a different condition often confused with this one. Kept so analysts can exclude it on purpose.", False),
    ("companion: the drug-cause code that should accompany a condition code (ICD-10-CM T45.AX5- and T45.1X5-).", False),
    ("exclude: not relevant. Excluded rows are dropped from the published table.", False),
    ("", False),
    ("Confidence", True),
    ("High: names the condition and is used for it in practice.", False),
    ("Moderate: right organ and process, but also used for non-immune causes.", False),
    ("Low: supporting or non-specific.", False),
    ("", False),
    ("Suggested order: Categories sheet first (keep, merge, or drop each category), then CTCAE (short), then ICD-10-CM and FAERS one category at a time. Use the filter arrow on category_id.", True),
    ("Rows with proposed_role = review have no suggestion; they need your call (accept does not work on them).", False),
    ("", False),
    ("Safe to re-run: if Claude rebuilds this file, your yellow columns are carried over automatically.", False),
]


def _sheet(wb, title, df, widths=None, author_cols=(), dv_specs=()):
    ws = wb.create_sheet(title)
    cols = list(df.columns)
    ws.append(cols)
    for row in df.itertuples(index=False):
        ws.append(["" if (v is None or (isinstance(v, float) and v != v)) else v for v in row])
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = BOLD
        cell.fill = YELLOW if c in author_cols else GREY if c.startswith("proposed_") else PatternFill()
        w = (widths or {}).get(c, min(45, max(10, len(c) + 2)))
        ws.column_dimensions[get_column_letter(j)].width = w
        if c in author_cols:
            for i in range(2, len(df) + 2):
                ws.cell(row=i, column=j).fill = YELLOW
    for col, formula in dv_specs:
        if col in cols:
            j = cols.index(col) + 1
            dv = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{max(2, len(df) + 1)}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    return ws


def write(path, cats, parts, of):
    wb = Workbook()
    ws = wb.active
    ws.title = "How to use"
    ws.column_dimensions["A"].width = 140
    for text, bold in HOW_TO:
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=bold, size=13 if ws.max_row == 1 else 11)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)
    ws.append([""])
    ws.append([f"FAERS counts: openFDA data last updated {of.get('openfda_last_updated')}, accessed {of.get('accessed')}; "
               f"{of['ici_report_total']:,} reports name a checkpoint inhibitor, of {of['all_report_total']:,} total. "
               "ici_prr compares how often a term appears in those reports versus all other reports (above 1 means over-represented)."])
    ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)

    _sheet(wb, "Categories", cats, widths={"label": 45, "lookalikes": 90, "author_note": 40},
           author_cols=("author_keep", "author_note"), dv_specs=[("author_keep", "keep,merge,drop")])

    author = ("author_role", "author_confidence", "author_rationale", "author_lookalike_note", "decided_by", "decided_on")
    dvs = [("author_role", ROLES), ("author_confidence", CONF)]
    w = {"category_label": 30, "source_term": 55, "proposed_rationale": 60, "author_rationale": 45,
         "author_lookalike_note": 35, "grade_1": 30, "grade_2": 30, "grade_3": 30, "grade_4": 30, "grade_5": 12}
    for name, df in parts.items():
        _sheet(wb, name, df, widths=w, author_cols=author, dv_specs=dvs)
    wb.save(path)
